# Confere o .xlsx gerado pelo app com dois programas de fora:
#  1) o LibreOffice recalcula todas as fórmulas: nenhuma pode dar erro, e o valor que o arquivo já traz
#     (o que a contadora vê ao abrir) tem que ser igual ao recalculado;
#  2) o validador oficial da Microsoft (Open XML SDK, em oxv/) confere a estrutura do arquivo para o
#     Office 2007, o Office 2016 e o Microsoft 365.
# Uso: python3 chaveiro/testes/valida_excel.py arquivo.xlsx
# A última linha é um JSON {"libreoffice": ..., "microsoft": ..., "falhas": n}.
# Programa que não está instalado aparece como "pulado" (instale com .claude/hooks/session-start.sh).
import json, os, re, shutil, subprocess, sys, tempfile
from pathlib import Path

FERR = Path(os.environ.get('CHAVEIRO_FERRAMENTAS') or Path.home() / '.cache' / 'chaveiro-ferramentas')
# perfil do LibreOffice que manda recalcular tudo ao abrir um .xlsx (0 = sempre)
XCU = '''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
'''
ERROS = ('#NAME?', '#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!', 'Err:')


def numero(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def libreoffice(arq):
    if not shutil.which('soffice'):
        return 'pulado'
    import openpyxl
    with tempfile.TemporaryDirectory(prefix='chaveiro-lo-') as tmp:
        perfil = Path(tmp) / 'perfil'
        (perfil / 'user').mkdir(parents=True)
        (perfil / 'user' / 'registrymodifications.xcu').write_text(XCU, encoding='utf-8')
        saida = Path(tmp) / 'saida'
        saida.mkdir()
        r = subprocess.run(['soffice', '-env:UserInstallation=' + perfil.as_uri(), '--headless', '--norestore',
                            '--convert-to', 'xlsx', '--outdir', str(saida), arq],
                           env=dict(os.environ, SAL_USE_VCLPLUGIN='svp'), capture_output=True, text=True, timeout=240)
        novo = saida / Path(arq).name
        if not novo.exists():
            return {'erro': 'o LibreOffice não abriu o arquivo: ' + (r.stderr or r.stdout)[-300:]}
        formulas = openpyxl.load_workbook(arq)
        cache = openpyxl.load_workbook(arq, data_only=True)
        recalc = openpyxl.load_workbook(novo, data_only=True)
        n, erros, difs = 0, [], []
        for ws in formulas:
            for linha in ws.iter_rows():
                for c in linha:
                    if not (isinstance(c.value, str) and c.value.startswith('=')):
                        continue
                    n += 1
                    a, b = cache[ws.title][c.coordinate].value, recalc[ws.title][c.coordinate].value
                    onde = f'{ws.title}!{c.coordinate} {c.value}'
                    if isinstance(b, str) and b.startswith(ERROS):
                        erros.append(f'{onde} deu {b}')
                    elif not (abs(float(a) - float(b)) < 0.006 if numero(a) and numero(b) else a == b or (a in (None, '') and b in (None, ''))):
                        difs.append(f'{onde}: no arquivo {a!r}, recalculado {b!r}')
        return {'formulas': n, 'erros': len(erros), 'diferentes': len(difs), 'detalhes': (erros + difs)[:10]}


def microsoft(arq):
    dll = FERR / 'oxv' / 'oxv.dll'
    if not (shutil.which('dotnet') and dll.exists()):
        return 'pulado'
    r = subprocess.run(['dotnet', str(dll), arq], capture_output=True, text=True, timeout=240)
    versoes = re.findall(r'\[(\w+)\]: (\d+) erro', r.stdout)
    if r.returncode not in (0, 1) or not versoes:
        return {'erro': 'o validador não rodou: ' + (r.stderr or r.stdout)[-300:]}
    return {'versoes': [v for v, _ in versoes], 'erros': sum(int(n) for _, n in versoes),
            'detalhes': [l.strip() for l in r.stdout.splitlines() if l.startswith('   ')][:10]}


if __name__ == '__main__':
    arq = os.path.abspath(sys.argv[1])
    res = {'libreoffice': libreoffice(arq), 'microsoft': microsoft(arq)}
    falhas = 0
    for nome, r in (('LibreOffice', res['libreoffice']), ('Microsoft', res['microsoft'])):
        if r == 'pulado':
            print(f'PULADO: {nome} não está instalado (rode .claude/hooks/session-start.sh)')
        elif 'erro' in r:
            falhas += 1
            print(f'{nome}: {r["erro"]}')
        else:
            falhas += r['erros'] + r.get('diferentes', 0)
            print(f'{nome}: {r}')
    res['falhas'] = falhas
    print(json.dumps(res, ensure_ascii=False))
    sys.exit(1 if falhas else 0)
